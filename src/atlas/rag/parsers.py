"""Document parsers for PDF, TXT, Markdown, DOCX and HTML."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from atlas.security.html import sanitise_html


@dataclass(slots=True)
class ParsedPage:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    filename: str
    content_type: str
    pages: list[ParsedPage] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text.strip())


class ParseError(ValueError):
    """Raised when a document cannot be parsed."""


def parse_document(path: Path, content_type: str | None = None) -> ParsedDocument:
    suffix = path.suffix.lower()
    data = path.read_bytes()
    try:
        if suffix == ".pdf":
            return _parse_pdf(path.name, data)
        if suffix in {".txt", ".md"}:
            return _parse_text(path.name, data, "text/markdown" if suffix == ".md" else "text/plain")
        if suffix == ".docx":
            return _parse_docx(path.name, data)
        if suffix in {".html", ".htm"}:
            return _parse_html(path.name, data)
    except Exception as exc:
        raise ParseError(f"failed to parse {path.name}: {exc}") from exc
    raise ParseError(f"unsupported format: {suffix}")


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    tmp_suffix = Path(filename).suffix.lower() or ".txt"
    # Parse from in-memory without requiring a durable temp file for text formats.
    if tmp_suffix in {".txt", ".md"}:
        return _parse_text(filename, data, "text/markdown" if tmp_suffix == ".md" else "text/plain")
    if tmp_suffix in {".html", ".htm"}:
        return _parse_html(filename, data)
    if tmp_suffix == ".pdf":
        return _parse_pdf(filename, data)
    if tmp_suffix == ".docx":
        return _parse_docx(filename, data)
    raise ParseError(f"unsupported format: {tmp_suffix}")


def _parse_text(filename: str, data: bytes, content_type: str) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    return ParsedDocument(
        filename=filename,
        content_type=content_type,
        pages=[ParsedPage(text=text, page=1, section="body")],
    )


def _parse_html(filename: str, data: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    raw = data.decode("utf-8", errors="replace")
    soup = BeautifulSoup(sanitise_html(raw), "lxml")
    text = soup.get_text("\n")
    return ParsedDocument(
        filename=filename,
        content_type="text/html",
        pages=[ParsedPage(text=text, page=1, section="body")],
    )


def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = [
        ParsedPage(text=page.extract_text() or "", page=index + 1, section=f"page-{index + 1}")
        for index, page in enumerate(reader.pages)
    ]
    if not any(page.text.strip() for page in pages):
        raise ParseError("PDF contained no extractable text")
    return ParsedDocument(filename=filename, content_type="application/pdf", pages=pages)


def _parse_docx(filename: str, data: bytes) -> ParsedDocument:
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(data))
    paragraphs = [para.text for para in document.paragraphs if para.text.strip()]
    text = "\n".join(paragraphs)
    if not text.strip():
        raise ParseError("DOCX contained no extractable text")
    return ParsedDocument(
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        pages=[ParsedPage(text=text, page=1, section="body")],
    )
